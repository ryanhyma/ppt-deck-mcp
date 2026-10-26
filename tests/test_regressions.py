"""Regression tests for every defect found while building the deck pipeline.

Run:  .venv/bin/python -m unittest discover -s tests -v

Each test names the failure it pins down, so a future change that reintroduces it fails here
instead of shipping a good-looking wrong deck.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import sys
import tempfile
import threading
import time
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent

# Environment comes from one place, before any server module is imported.
sys.path.insert(0, str(HERE))
from _env import TMP as _TMP  # noqa: E402

_TEMPLATES = ROOT / "templates"

sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "upstream"))
sys.path.insert(0, str(HERE))

import server  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

import deck_preview  # noqa: E402
import render_cache  # noqa: E402
import render_warm  # noqa: E402
from deck_tools import _outline, _pick_layouts  # noqa: E402

TOOLS = {name: tool.fn for name, tool in server.app._tool_manager._tools.items()}


def build(**kwargs):
    base = {"template": "corporate-2slide.pptx", "title": "Regression deck", "slides": []}
    base.update(kwargs)
    return TOOLS["build_deck"](**base)


def file_of(result: dict) -> str:
    """Tests inspect the real .pptx; the MCP surface only returns a deck id."""
    from deck_preview import resolve_ref

    return str(resolve_ref(result["deck_id"]))


def assert_signed_download(case, url):
    """A deck download must be a signed, expiring URL — not an open path."""
    from urllib.parse import parse_qs, urlsplit

    parts = urlsplit(url)
    case.assertTrue(parts.path.endswith("/deck.pptx"), url)
    query = parse_qs(parts.query)
    case.assertIn("exp", query)
    case.assertIn("sig", query)
    case.assertGreaterEqual(len(query["sig"][0]), 32)


def bullets(n, length=40):
    return [f"Point {i} " + "x" * max(0, length - 8) for i in range(1, n + 1)]


def body_of(prs, index):
    for ph in prs.slides[index].placeholders:
        if ph.placeholder_format.idx == 1:
            return ph
    return None


class LayoutRegressions(unittest.TestCase):
    def test_content_layout_is_title_and_content(self):
        """Was: content slides used layout 10 ("Content with Caption")."""
        prs = Presentation(str(_TEMPLATES / "corporate-2slide.pptx"))
        title_idx, content_idx = _pick_layouts(prs)
        self.assertEqual(title_idx, 0)
        self.assertEqual(content_idx, 1)

    def test_narrow_bullet_column_does_not_grow_over_the_title(self):
        """Was: with a chart beside it, the narrowed body placeholder grew upward into the title."""
        out = build(slides=[{
            "title": "Overlap check",
            "bullets": bullets(3),
            "chart": {"type": "column", "categories": ["A", "B"], "series": [{"name": "S", "values": [1, 2]}]},
        }])
        prs = Presentation(file_of(out))
        slide = prs.slides[1]
        title = next(ph for ph in slide.placeholders if ph.placeholder_format.idx == 0)
        body = body_of(prs, 1)
        self.assertGreaterEqual(body.top, title.top + title.height - Inches(0.15))
        self.assertEqual(body.top, Inches(1.55))
        self.assertEqual(body.width, Inches(5.4))

    def test_visual_left_puts_the_chart_left_of_the_text(self):
        out = build(slides=[{
            "title": "Left visual",
            "bullets": bullets(2),
            "layout": "image-left",
            "chart": {"type": "bar", "categories": ["A"], "series": [{"name": "S", "values": [3]}]},
        }])
        prs = Presentation(file_of(out))
        slide = prs.slides[1]
        chart = next(sh for sh in slide.shapes if sh.has_chart)
        body = body_of(prs, 1)
        self.assertLess(chart.left, body.left)

    def test_visual_full_drops_the_bullet_placeholder(self):
        out = build(slides=[{
            "title": "No bullets",
            "bullets": bullets(2),
            "layout": "visual-full",
            "chart": {"type": "line", "categories": ["A", "B"], "series": [{"name": "S", "values": [1, 2]}]},
        }])
        prs = Presentation(file_of(out))
        self.assertIsNone(body_of(prs, 1))


class ValidationRegressions(unittest.TestCase):
    def test_table_without_rows_is_rejected(self):
        """Was: silently skipped, producing a deck whose 'Pricing Comparison' slide had no table."""
        with self.assertRaises(ValueError) as ctx:
            build(slides=[{"title": "Table", "table": {"headers": ["a", "b"], "rows": []}}])
        self.assertIn("rows", str(ctx.exception))

    def test_chart_without_series_is_rejected(self):
        with self.assertRaises(ValueError) as ctx:
            build(slides=[{"title": "Chart", "chart": {"categories": ["a"]}}])
        self.assertIn("series", str(ctx.exception))

    def test_unknown_template_fails_loudly(self):
        with self.assertRaises(FileNotFoundError):
            build(template="corporate", slides=[])


class OutlineRegressions(unittest.TestCase):
    def test_outline_ignores_decorative_chrome(self):
        """Was: the footer wordmark came back as the title-slide subtitle and got rebuilt as content."""
        out = build(title="Deck", subtitle="", slides=[{"title": "Body", "bullets": bullets(2)}])
        data = _outline(Path(file_of(out)))
        self.assertEqual(data["title"], "Deck")
        self.assertEqual(data["subtitle"], "")
        self.assertNotIn("Confidential", data["subtitle"])
        self.assertEqual(len(data["slides"][1]["bullets"]), 2)

    def test_outline_round_trips_chart_data(self):
        out = build(slides=[{
            "title": "Chart",
            "chart": {"type": "column", "categories": ["A", "B"], "series": [{"name": "2026", "values": [10, 20]}]},
        }])
        data = _outline(Path(file_of(out)))
        chart = data["slides"][1]["chart"]
        self.assertEqual(chart["categories"], ["A", "B"])
        self.assertEqual(chart["series"][0]["values"], [10.0, 20.0])


class StyleRegressions(unittest.TestCase):
    def test_title_bold_does_not_bold_the_body(self):
        """Was: one `bold` flag hit every placeholder, so "bold the title" bolded the bullets too."""
        out = build(slides=[{"title": "T", "bullets": bullets(2)}])
        deck = file_of(out)
        TOOLS["style_slide"](file=deck, slide_index=2, title_bold=True)
        prs = Presentation(deck)
        slide = prs.slides[1]
        title = next(ph for ph in slide.placeholders if ph.placeholder_format.idx == 0)
        body = body_of(prs, 1)
        self.assertTrue(title.text_frame.paragraphs[0].runs[0].font.bold)
        self.assertFalse(bool(body.text_frame.paragraphs[0].runs[0].font.bold))

    def test_styling_survives_a_content_rebuild_with_base(self):
        """Was: a content edit rebuilt from the outline and wiped the pane's per-slide styling."""
        out = build(slides=[{"title": "One", "bullets": bullets(2)}, {"title": "Two", "bullets": bullets(2)}])
        deck = file_of(out)
        TOOLS["style_slide"](file=deck, slide_index=2, title_size=40, title_color="#C4703C", font="Georgia")

        rebuilt = build(
            title="Regression deck",
            slides=[{"title": "One", "bullets": bullets(2)}, {"title": "Two edited", "bullets": bullets(3)}],
            base=deck,
            output_name="rebuilt",
        )
        prs = Presentation(file_of(rebuilt))
        title = next(ph for ph in prs.slides[1].placeholders if ph.placeholder_format.idx == 0)
        run = title.text_frame.paragraphs[0].runs[0]
        self.assertEqual(run.font.size, Pt(40))
        self.assertEqual(str(run.font.color.rgb), "C4703C")
        self.assertEqual(run.font.name, "Georgia")
        self.assertEqual(len(body_of(prs, 2).text_frame.text.split("\n")), 3)

    def test_style_text_emphasises_only_the_matching_run(self):
        out = build(slides=[{"title": "T", "bullets": ["Bundling raises retention", "Other point"]}])
        deck = file_of(out)
        result = TOOLS["style_text"](file=deck, slide_index=2, phrase="retention", bold=True, color="#C4703C")
        self.assertEqual(result["runs_changed"], 1)
        prs = Presentation(deck)
        runs = [r for p in body_of(prs, 1).text_frame.paragraphs for r in p.runs]
        target = next(r for r in runs if "retention" in r.text)
        other = next(r for r in runs if "Other" in r.text)
        self.assertTrue(target.font.bold)
        self.assertEqual(str(target.font.color.rgb), "C4703C")
        self.assertFalse(bool(other.font.bold))

    def test_style_text_fails_loudly_when_phrase_is_absent(self):
        """A missing phrase must not be a silent no-op."""
        out = build(slides=[{"title": "T", "bullets": ["Something else"]}])
        with self.assertRaises(ValueError) as ctx:
            TOOLS["style_text"](file=file_of(out), slide_index=2, phrase="nonexistent", bold=True)
        self.assertIn("not found", str(ctx.exception))

    def test_dense_slides_shrink_bullets_and_warn(self):
        out = build(slides=[{"title": "Dense", "bullets": bullets(6, length=70)}])
        self.assertTrue(out["warnings"], "expected a shrink warning")
        prs = Presentation(file_of(out))
        run = body_of(prs, 1).text_frame.paragraphs[0].runs[0]
        self.assertEqual(run.font.size, Pt(17))

    def test_short_six_bullets_keep_full_size(self):
        """6 short bullets at 17 pt looked anaemic — the shrink is length-gated."""
        out = build(slides=[{"title": "Short", "bullets": ["One", "Two", "Three", "Four", "Five", "Six"]}])
        self.assertEqual(out["warnings"], [])
        prs = Presentation(file_of(out))
        self.assertEqual(body_of(prs, 1).text_frame.paragraphs[0].runs[0].font.size, Pt(20))

    def test_chart_colours_override_the_palette(self):
        out = build(slides=[{
            "title": "Colour",
            "chart": {"type": "column", "categories": ["A"], "colors": ["#C4703C"],
                      "series": [{"name": "S", "values": [1]}]},
        }])
        prs = Presentation(file_of(out))
        chart = next(sh.chart for sh in prs.slides[1].shapes if sh.has_chart)
        self.assertEqual(str(chart.plots[0].series[0].format.fill.fore_color.rgb), "C4703C")

    def test_dark_template_uses_light_body_text(self):
        """Was: layout-level styling never reached empty placeholders -> black text on a dark deck."""
        out = build(template="midnight.pptx", slides=[{"title": "Dark", "bullets": bullets(1)}])
        prs = Presentation(file_of(out))
        run = body_of(prs, 1).text_frame.paragraphs[0].runs[0]
        self.assertEqual(str(run.font.color.rgb), "C3CDDB")

    def test_output_name_does_not_gain_a_pptx_suffix(self):
        out = build(output_name="quarterly.pptx", slides=[])
        self.assertNotIn("pptx-", Path(file_of(out)).stem)


class ImportRegressions(unittest.TestCase):
    def test_missing_absolute_path_reports_clearly(self):
        """Was: pathlib raised 'Non-relative patterns are unsupported' for absolute paths."""
        with self.assertRaises(FileNotFoundError) as ctx:
            TOOLS["import_deck"](file="/home/nobody/nope.pptx")
        self.assertIn("no file at", str(ctx.exception))

    def test_bare_filename_resolves_inside_uploads(self):
        src = _TMP / "uploads" / "user" / "Q3 Review.pptx"
        shutil.copyfile(_TEMPLATES / "corporate-2slide.pptx", src)
        result = TOOLS["import_deck"](file="Q3 Review.pptx")
        self.assertEqual(result["imported_name"], src.name)
        self.assertEqual(result["slide_count"], 2)

    def test_legacy_pot_is_rejected(self):
        legacy = _TMP / "uploads" / "user" / "old.pot"
        legacy.write_bytes(b"not a zip")
        with self.assertRaises(ValueError):
            TOOLS["import_deck"](file="old.pot")


class RendererBackendRegressions(unittest.TestCase):
    """The preview renderer is optional and swappable; decks must survive every choice."""

    def test_explicit_backends_are_honoured(self):
        from deck_preview import _pick_renderer

        original = os.environ.get("PPT_RENDERER")
        try:
            for value, expected in (("none", "none"), ("soffice", "soffice"),
                                    ("libreoffice", "soffice"), ("powerpoint", "powerpoint"),
                                    ("com", "powerpoint"), ("aspose", "aspose")):
                os.environ["PPT_RENDERER"] = value
                self.assertEqual(_pick_renderer(), expected, value)
        finally:
            if original is None:
                os.environ.pop("PPT_RENDERER", None)
            else:
                os.environ["PPT_RENDERER"] = original

    def test_auto_picks_something_usable(self):
        from deck_preview import _pick_renderer

        original = os.environ.pop("PPT_RENDERER", None)
        try:
            choice = _pick_renderer()
            self.assertIn(choice, {"soffice", "powerpoint", "none"})
            if os.name != "nt" and not shutil.which("soffice"):
                self.assertEqual(choice, "none")
        finally:
            if original is not None:
                os.environ["PPT_RENDERER"] = original

    def test_aspose_backend_reports_missing_dependency(self):
        """aspose-slides is optional; without it the tool degrades, it does not crash."""
        from deck_preview import RendererUnavailable, _render_aspose

        try:
            import aspose.slides  # noqa: F401
            self.skipTest("aspose-slides installed in this environment")
        except ImportError:
            pass
        with self.assertRaises(RendererUnavailable) as ctx:
            _render_aspose(Path(_TEMPLATES / "corporate-2slide.pptx"), _TMP / "out")
        self.assertIn("aspose", str(ctx.exception).lower())

    def test_config_file_supplies_soffice_bin(self):
        """Windows setup is a JSON file, not a pile of setx commands."""
        import deck_preview

        cfg = _TMP / "cfg.json"
        cfg.write_text(json.dumps({"soffice_bin": "C:\\tools\\lo\\soffice.exe"}))
        original = (os.environ.get("PPT_CONFIG"), os.environ.pop("SOFFICE_BIN", None))
        os.environ["PPT_CONFIG"] = str(cfg)
        try:
            self.assertEqual(deck_preview._soffice_bin(), r"C:\tools\lo\soffice.exe")
            os.environ["SOFFICE_BIN"] = "/env/wins/soffice"
            self.assertEqual(deck_preview._soffice_bin(), "/env/wins/soffice", "env must win")
        finally:
            os.environ.pop("SOFFICE_BIN", None)
            if original[0] is None:
                os.environ.pop("PPT_CONFIG", None)
            else:
                os.environ["PPT_CONFIG"] = original[0]
            if original[1] is not None:
                os.environ["SOFFICE_BIN"] = original[1]

    def test_windows_hints_point_at_portable_locations(self):
        from deck_preview import _WINDOWS_SOFFICE_HINTS

        self.assertTrue(any("LibreOfficePortable" in hint for hint in _WINDOWS_SOFFICE_HINTS))
        self.assertTrue(all(hint.lower().endswith("soffice.exe") for hint in _WINDOWS_SOFFICE_HINTS))

    def test_pdf_rasterises_without_poppler(self):
        """The PDF->PNG step must work from a pip wheel: Windows has no pdftoppm."""
        import pymupdf

        from deck_preview import _pdf_to_pngs

        pdf = _TMP / "sample.pdf"
        doc = pymupdf.open()
        page = doc.new_page(width=320, height=180)
        page.insert_text((20, 60), "slide")
        doc.save(str(pdf))
        doc.close()

        outdir = _TMP / "raster"
        outdir.mkdir(exist_ok=True)
        count = _pdf_to_pngs(pdf, outdir, dpi=72)
        self.assertEqual(count, 1)
        png = outdir / "slide-1.png"
        self.assertTrue(png.is_file())
        self.assertGreater(png.stat().st_size, 200)

    def test_portable_soffice_path_is_used(self):
        """SOFFICE_BIN points at an unzipped LibreOffice, so no install/PATH change is needed."""
        from deck_preview import RendererUnavailable, _render_soffice

        original = os.environ.get("SOFFICE_BIN")
        os.environ["SOFFICE_BIN"] = "/nonexistent/soffice"
        try:
            with self.assertRaises(RendererUnavailable) as ctx:
                _render_soffice(Path(_TEMPLATES / "corporate-2slide.pptx"), _TMP / "out")
            self.assertIn("/nonexistent/soffice", str(ctx.exception))
        finally:
            if original is None:
                os.environ.pop("SOFFICE_BIN", None)
            else:
                os.environ["SOFFICE_BIN"] = original

    def test_disabled_renderer_degrades_instead_of_failing(self):
        from deck_preview import build_preview

        deck = file_of(build(slides=[{"title": "T", "bullets": bullets(1)}]))
        original = os.environ.get("PPT_RENDERER")
        os.environ["PPT_RENDERER"] = "none"
        try:
            data = build_preview(deck)
        finally:
            if original is None:
                os.environ.pop("PPT_RENDERER", None)
            else:
                os.environ["PPT_RENDERER"] = original
        self.assertEqual(data["slides"], [])
        self.assertIn("renderer", data["preview_error"].lower())
        self.assertEqual(data["slide_count"], 2)
        assert_signed_download(self, data["download"])


class PreviewRegressions(unittest.TestCase):
    """A missing renderer must not fail a styling operation (it did on the GitHub runner)."""

    def _break_renderer(self):
        import deck_preview

        original = deck_preview._render
        deck_preview._render = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("soffice not found"))
        # The warm renderer is a second path to a picture: breaking only the cold chain no longer
        # means "no renderer at all", so it has to be taken out of the picture too.
        original_warm = deck_preview.warm_enabled
        deck_preview.warm_enabled = lambda: False
        return deck_preview, (original, original_warm)

    def test_build_preview_reports_the_error_instead_of_raising(self):
        deck = file_of(build(slides=[{"title": "T", "bullets": bullets(1)}]))
        module, (original, original_warm) = self._break_renderer()
        try:
            from deck_preview import build_preview

            data = build_preview(deck)
        finally:
            module._render = original
            module.warm_enabled = original_warm
        self.assertEqual(data["slides"], [])
        self.assertIn("soffice not found", data["preview_error"])
        # The deck's size must not depend on whether a renderer exists.
        self.assertEqual(data["slide_count"], 2)
        self.assertEqual(data["rendered_slides"], 0)
        assert_signed_download(self, data["download"])

    def test_style_slide_still_applies_without_a_renderer(self):
        out = build(slides=[{"title": "T", "bullets": bullets(1)}])
        module, (original, original_warm) = self._break_renderer()
        try:
            result = TOOLS["style_slide"](file=file_of(out), slide_index=2, title_size=30)
        finally:
            module._render = original
            module.warm_enabled = original_warm
        self.assertEqual(result["slides"], [])
        prs = Presentation(file_of(out))
        title = next(ph for ph in prs.slides[1].placeholders if ph.placeholder_format.idx == 0)
        self.assertEqual(title.text_frame.paragraphs[0].runs[0].font.size, Pt(30))


class VerifyRegressions(unittest.TestCase):
    """verify_deck is the safety net for the failure class the model keeps repeating."""

    def test_flags_missing_table_and_wrong_slide_count(self):
        deck = file_of(build(slides=[{"title": "One", "bullets": bullets(2)}]))
        result = TOOLS["verify_deck"](file=deck, expected_slides=3, table=True, chart=False, image=False)
        self.assertFalse(result["ok"])
        self.assertTrue(any("slide count" in p for p in result["problems"]))
        self.assertTrue(any("no table" in p for p in result["problems"]))
        self.assertEqual(result["slide_count"], 2)

    def test_passes_when_the_deck_matches_the_ask(self):
        built = build(slides=[
            {"title": "One", "table": {"headers": ["a", "b"], "rows": [["1", "2"]]}},
            {"title": "Two", "chart": {"type": "column", "categories": ["A"],
                                       "series": [{"name": "S", "values": [1]}]}},
        ])
        deck = file_of(built)
        result = TOOLS["verify_deck"](file=deck, expected_slides=3, table=True, chart=True)
        self.assertTrue(result["ok"], result["problems"])
        self.assertEqual((result["tables"], result["charts"]), (1, 1))

    def test_missing_file_is_a_clear_error(self):
        with self.assertRaises(FileNotFoundError):
            TOOLS["verify_deck"](file=str(_TMP / "nope.pptx"))


class OneVisualPerSlideRegressions(unittest.TestCase):
    """Was: chart+table on one slide silently dropped the table (root cause of "table missing")."""

    def test_two_visuals_on_one_slide_is_rejected(self):
        with self.assertRaises(ValueError) as ctx:
            build(slides=[{
                "title": "Both",
                "chart": {"type": "column", "categories": ["A"], "series": [{"name": "S", "values": [1]}]},
                "table": {"headers": ["a"], "rows": [["1"]]},
            }])
        self.assertIn("one visual per slide", str(ctx.exception))

    def test_table_and_chart_on_separate_slides_both_survive(self):
        out = build(slides=[
            {"title": "Table slide", "table": {"headers": ["a", "b"], "rows": [["1", "2"]]}},
            {"title": "Chart slide", "chart": {"type": "pie", "categories": ["A", "B"],
                                               "series": [{"name": "S", "values": [1, 2]}]}},
        ])
        check = TOOLS["verify_deck"](file=file_of(out), expected_slides=3, table=True, chart=True)
        self.assertTrue(check["ok"], check["problems"])


class ConcurrencyRegressions(unittest.TestCase):
    """Rendering and building must survive several users at once."""

    def test_parallel_builds_do_not_collide(self):
        from concurrent.futures import ThreadPoolExecutor

        def make(i):
            return file_of(build(title=f"Deck {i}", slides=[{"title": f"S{i}", "bullets": bullets(2)}],
                                 output_name="parallel"))

        with ThreadPoolExecutor(max_workers=6) as pool:
            files = list(pool.map(make, range(6)))
        self.assertEqual(len(set(files)), 6, "output names must be unique")
        for f in files:
            self.assertEqual(len(Presentation(f).slides), 2)

    def test_same_deck_rendered_concurrently_is_consistent(self):
        from concurrent.futures import ThreadPoolExecutor

        from deck_preview import build_preview

        built = build(slides=[{"title": "A", "bullets": bullets(1)},
                              {"title": "B", "bullets": bullets(1)}])
        deck_id = built["deck_id"]
        with ThreadPoolExecutor(max_workers=5) as pool:
            results = list(pool.map(lambda _: build_preview(deck_id), range(5)))
        for data in results:
            self.assertEqual(data["slide_count"], 3)
            assert_signed_download(self, data["download"])
        # the surviving preview must be complete and readable
        first = results[-1]
        self.assertEqual(first["slide_count"], 3)

    def test_concurrent_style_and_preview_keep_the_deck_valid(self):
        from concurrent.futures import ThreadPoolExecutor

        from deck_preview import build_preview

        built = build(slides=[{"title": "A", "bullets": bullets(2)}])
        deck = file_of(built)          # to inspect the artefact afterwards
        deck_ref = built["deck_id"]    # what a client actually holds

        def work(i):
            if i % 2:
                return TOOLS["style_slide"](file=deck_ref, slide_index=2, body_size=18 + i)
            return build_preview(deck_ref)

        with ThreadPoolExecutor(max_workers=4) as pool:
            list(pool.map(work, range(8)))
        # the file must still open and be complete after interleaved saves and renders
        self.assertEqual(len(Presentation(deck).slides), 2)


class HarnessRegressions(unittest.TestCase):
    """The suite's judge must not turn process noise into failures, or hide real ones."""

    CASE = {"id": "x", "category": "t", "prompt": "p",
            "expect": {"slides": 5, "chart": False, "table": False, "image": False,
                       "template": "corporate-2slide.pptx"}}
    GOT = {"slides": 5, "charts": 0, "tables": 0, "pictures": 0,
           "max_bullet_chars": 10, "bullet_overflow_slides": []}

    def judge(self, got, calls=(), error=None):
        from run_suite import judge
        return judge(self.CASE, got, list(calls), error)

    def test_tool_error_is_a_warning_not_a_failure(self):
        calls = [{"name": "build_deck_mcp_ppt", "output": "Error executing tool build_deck: boom"},
                 {"name": "preview_deck_mcp_ppt", "output": "{}"}]
        ok, fails, warnings = self.judge(self.GOT, calls)
        self.assertTrue(ok, fails)
        self.assertTrue(any("tool errors" in w for w in warnings))

    def test_missing_preview_is_a_warning(self):
        calls = [{"name": "build_deck_mcp_ppt", "output": "{}"}]
        ok, fails, warnings = self.judge(self.GOT, calls)
        self.assertTrue(ok)
        self.assertIn("preview_deck not called", warnings)

    def test_slide_count_mismatch_fails(self):
        ok, fails, _ = self.judge({**self.GOT, "slides": 6})
        self.assertFalse(ok)
        self.assertTrue(any("slides" in f for f in fails))

    def test_no_deck_fails(self):
        ok, fails, _ = self.judge(None, [])
        self.assertFalse(ok)
        self.assertIn("no deck produced", fails)


class PublicBaseRegressions(unittest.TestCase):
    """PPT_PUBLIC_BASE must be a hostname, never an IP address literal.

    LibreChat's MCP App bridge refuses to open host-opened links whose host is an IP
    literal, so a numeric base silently disables the Deck Studio "Download .pptx" button.
    """

    def test_rejects_ipv4_literal(self):
        from deck_preview import public_base

        original = os.environ.get("PPT_PUBLIC_BASE")
        try:
            os.environ["PPT_PUBLIC_BASE"] = "http://172.17.0.1:8931"
            with self.assertRaisesRegex(RuntimeError, "IP address"):
                public_base()
        finally:
            if original is None:
                os.environ.pop("PPT_PUBLIC_BASE", None)
            else:
                os.environ["PPT_PUBLIC_BASE"] = original

    def test_rejects_ipv6_literal(self):
        from deck_preview import public_base

        original = os.environ.get("PPT_PUBLIC_BASE")
        try:
            os.environ["PPT_PUBLIC_BASE"] = "http://[::1]:8931"
            with self.assertRaises(RuntimeError):
                public_base()
        finally:
            if original is None:
                os.environ.pop("PPT_PUBLIC_BASE", None)
            else:
                os.environ["PPT_PUBLIC_BASE"] = original

    def test_accepts_hostname(self):
        from deck_preview import public_base

        original = os.environ.get("PPT_PUBLIC_BASE")
        try:
            os.environ["PPT_PUBLIC_BASE"] = "http://decks.local:8931"
            self.assertEqual(public_base(), "http://decks.local:8931")
        finally:
            if original is None:
                os.environ.pop("PPT_PUBLIC_BASE", None)
            else:
                os.environ["PPT_PUBLIC_BASE"] = original

    def test_default_is_hostname(self):
        from deck_preview import public_base

        original = os.environ.get("PPT_PUBLIC_BASE")
        try:
            os.environ.pop("PPT_PUBLIC_BASE", None)
            self.assertEqual(public_base(), "http://host.docker.internal:8931")
        finally:
            if original is None:
                os.environ.pop("PPT_PUBLIC_BASE", None)
            else:
                os.environ["PPT_PUBLIC_BASE"] = original


class ToolTitleRegressions(unittest.TestCase):
    """Tool ids stay ASCII, so the human label must ride in title/annotations.title.

    LibreChat agent configs reference ids like `preview_deck_mcp_ppt`; renaming a tool to make
    it readable would silently break them. Hosts that display a label read `annotations.title`
    (older spec) or the top-level `title` (mcp >= 1.30), so both are set.
    """

    def _titles(self):
        import asyncio

        tools = asyncio.run(server.app.list_tools())
        return {t.name: (t.title, t.annotations.title if t.annotations else None) for t in tools}

    def test_preview_tools_carry_a_bilingual_title(self):
        titles = self._titles()
        for name in ("preview_deck", "get_last_deck", "list_decks"):
            self.assertIn(name, titles, f"{name} disappeared from the tool catalog")
            top_level, annotation = titles[name]
            self.assertTrue(top_level, f"{name} has no top-level title")
            self.assertEqual(top_level, annotation, f"{name}: title and annotations.title disagree")
            self.assertIn("/", top_level, f"{name}: title should be bilingual (English / French)")
            self.assertRegex(top_level, r"[A-Za-z]", "title must stay ASCII-typable")

    def test_tool_ids_stay_ascii_identifiers(self):
        for name in self._titles():
            self.assertRegex(name, r"^[A-Za-z0-9_-]{1,64}$")


class DeckStudioAppAccessibility(unittest.TestCase):
    """The MCP App view must stay keyboard-operable and WCAG 2.2 AA shaped.

    The app is plain HTML/JS served from disk, so these are the only checks that run
    without a browser; the interactive half lives in tests/e2e/deck-studio.js.
    """

    @classmethod
    def setUpClass(cls):
        cls.html = (ROOT / "apps" / "deck_studio" / "index.html").read_text(encoding="utf-8")
        cls.js = (ROOT / "apps" / "deck_studio" / "app.js").read_text(encoding="utf-8")

    def test_language_is_canadian_english(self):
        self.assertIn('<html lang="en-CA">', self.html)

    def test_slide_viewer_is_a_modal_dialog(self):
        self.assertIn('<dialog id="viewer"', self.html)
        self.assertIn('aria-labelledby="viewer-title"', self.html)
        for control in ('viewer-prev', 'viewer-next', 'viewer-close'):
            self.assertIn(f'id="{control}"', self.html)
        # Arrow keys are advertised to assistive tech, not just implemented.
        self.assertIn('aria-keyshortcuts="ArrowLeft"', self.html)
        self.assertIn('aria-keyshortcuts="ArrowRight"', self.html)

    def test_viewer_handles_arrow_keys_and_escape(self):
        for key in ('ArrowRight', 'ArrowLeft', 'Home', 'End', 'Escape'):
            self.assertIn(key, self.js, f"{key} handling removed from the viewer")

    def test_status_is_announced_politely(self):
        self.assertIn('role="status"', self.html)
        self.assertIn('aria-live="polite"', self.html)

    def test_slide_controls_are_not_hover_only(self):
        # The old build revealed Style / Ask agent buttons with `figure:hover`, which made
        # them invisible to touch, keyboard-only and low-vision users.
        self.assertNotIn('figure:hover .edit', self.html)
        self.assertNotIn('figure:hover .stylebtn', self.html)
        self.assertIn('.card-actions', self.html)
        self.assertIn('"card-actions"', self.js)

    def test_every_slide_image_gets_a_text_alternative(self):
        self.assertIn('img.alt =', self.js)
        self.assertIn('viewSlide', self.js)
        # Decorative thumbnails must stay out of the accessibility tree.
        self.assertIn('img.alt = ""', self.js)

    def test_both_languages_define_every_ui_string(self):
        """A missing fr-CA string would silently fall back to English in a French UI."""
        keys = set(re.findall(r'data-i18n(?:-aria|-placeholder)?="([^"]+)"', self.html))
        keys |= set(re.findall(r'\bt\("([A-Za-z][A-Za-z0-9_]*)', self.js))
        self.assertGreater(len(keys), 20, "expected the UI strings to be externalised")
        missing = sorted(k for k in keys if self.js.count(f'{k}: "') < 2)
        self.assertEqual(missing, [], f"missing en/fr entries: {missing}")

    def test_language_switch_carries_its_own_lang_and_state(self):
        self.assertIn('id="lang-fr"', self.html)
        self.assertIn('id="lang-en"', self.html)
        self.assertIn('aria-pressed', self.html)
        # <html lang> must follow the switch, or screen readers keep the English voice.
        self.assertIn("document.documentElement.lang", self.js)
        self.assertIn('fr-CA', self.js)

    def test_french_strings_are_canadian(self):
        for phrase in ('Fermer', 'Diapositive', 'Télécharger', 'Annuler'):
            self.assertIn(phrase, self.js, f"missing French wording: {phrase}")

    def test_no_teal_button_below_contrast_minimum(self):
        # #2a9d8f on white is 3.3:1 and fails AA for body text; the UI uses #17766a (5.5:1).
        self.assertNotIn('--btn-bg: #2a9d8f', self.html)
        self.assertIn('--btn-bg: #17766a', self.html)

    def test_reduced_motion_is_honoured(self):
        self.assertIn('prefers-reduced-motion', self.html)


# ----------------------------------------------------------------- warm renderer + slide cache
class _FakeRenderer:
    """Stand-in for render_warm.WarmRenderer: records what would have been rendered."""

    def __init__(self, root: Path, fail: bool = False, width: int = 1600):
        self.width = width
        self.calls: list[tuple[str, list[int]]] = []
        self.fail = fail
        self._root = root

    def render(self, src, wanted, root):
        self.calls.append((Path(src).name, sorted(wanted)))
        if self.fail:
            raise RuntimeError("bridge died")
        out = {}
        for index, key in wanted.items():
            png = root / f"fake-{key}.png"
            png.parent.mkdir(parents=True, exist_ok=True)
            png.write_bytes(b"\x89PNG-fake-" + key[:8].encode())
            render_cache.store(root, key, png)
            out[index] = render_cache.path_for(root, key)
        return out

    def live(self):
        return True

    def stats(self):
        return {}


class SlideCacheRegressions(unittest.TestCase):
    """The cache exists so an unchanged slide is never rendered twice."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(dir=_TMP, prefix="cache-"))
        # Isolate the preview root per test: a shared cache would make these assertions depend on
        # which test ran first.
        self._saved_out = os.environ.get("PPT_OUTPUT_DIR")
        os.environ["PPT_OUTPUT_DIR"] = str(self.tmp / "out")
        self.deck = self.tmp / "deck.pptx"
        shutil.copyfile(_TEMPLATES / "corporate-2slide.pptx", self.deck)
        out = TOOLS["build_deck"](template="corporate-2slide.pptx", title="Cache deck",
                                  slides=[{"title": "One", "bullets": bullets(3)},
                                          {"title": "Two", "bullets": bullets(3)},
                                          {"title": "Three", "bullets": bullets(3)}])
        shutil.copyfile(file_of(out), self.deck)
        self.prs = Presentation(str(self.deck))
        self.slide_total = len(self.prs.slides)
        self.root = render_cache.cache_root(deck_preview.preview_root())
        self._saved = (deck_preview.render_warm.get_renderer, deck_preview.render_warm.engine_tag,
                       deck_preview._render)
        deck_preview.render_warm.engine_tag = lambda: "test-engine-1"

    def tearDown(self):
        deck_preview.render_warm.get_renderer, deck_preview.render_warm.engine_tag, \
            deck_preview._render = self._saved
        if self._saved_out is None:
            os.environ.pop("PPT_OUTPUT_DIR", None)
        else:
            os.environ["PPT_OUTPUT_DIR"] = self._saved_out
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _keys(self, prs=None):
        prs = prs or Presentation(str(self.deck))
        return render_cache.slide_keys(self.deck, [str(s.part.partname) for s in prs.slides],
                                       "test-engine-1", deck_preview.preview_width())

    def test_only_the_edited_slide_changes_its_key(self):
        """A per-slide key is the whole point; a deck-level key would re-render everything."""
        before = self._keys()
        TOOLS["style_slide"](file=str(self.deck), slide_index=2, title_size=44)
        after = self._keys()
        changed = [i for i, (a, b) in enumerate(zip(before, after), start=1) if a != b]
        self.assertEqual(changed, [2])

    def test_shared_parts_invalidate_every_slide(self):
        """Theme/master/layout edits change how all slides look, so they must miss."""
        before = self._keys()
        prs = Presentation(str(self.deck))
        for slide in prs.slides:  # touching the layout surface changes shared parts
            slide.shapes
        layout = prs.slides[0].slide_layout
        for ph in layout.placeholders:
            if ph.has_text_frame:
                ph.text_frame.text = ""
        prs.save(str(self.deck))
        after = self._keys()
        self.assertNotEqual(before, after)

    def test_resolution_is_part_of_the_key(self):
        shared = render_cache.shared_digest(self.deck)
        a = render_cache.slide_key(shared, "e", 1600, b"<slide/>")
        b = render_cache.slide_key(shared, "e", 1200, b"<slide/>")
        self.assertNotEqual(a, b, "a 1200px render must not be served for a 1600px request")

    def test_second_preview_renders_nothing(self):
        fake = _FakeRenderer(self.root)
        deck_preview.render_warm.get_renderer = lambda **kw: fake
        first = deck_preview._prepare_previews(self.deck, self.tmp / "a", Presentation(str(self.deck)))
        second = deck_preview._prepare_previews(self.deck, self.tmp / "b", Presentation(str(self.deck)))
        deck_renders = [call for call in fake.calls if call[0] == self.deck.name]
        self.assertEqual(len(deck_renders), 1, "the second preview must not reach the renderer")
        self.assertEqual(first["misses"], self.slide_total)
        self.assertEqual(second["misses"], 0)
        self.assertEqual(second["renderer"], "cache")
        self.assertEqual(second["hits"], self.slide_total)
        self.assertTrue((self.tmp / "b" / "slide-3.png").is_file())
        self.assertEqual((self.tmp / "a" / "slide-3.png").read_bytes(),
                         (self.tmp / "b" / "slide-3.png").read_bytes())

    def test_one_edited_slide_renders_one_slide(self):
        fake = _FakeRenderer(self.root)
        deck_preview.render_warm.get_renderer = lambda **kw: fake
        deck_preview._prepare_previews(self.deck, self.tmp / "a", Presentation(str(self.deck)))
        TOOLS["style_text"](file=str(self.deck), slide_index=3, phrase="Point", bold=True)
        deck_calls = [call for call in fake.calls if call[0] == self.deck.name]
        # The tool refreshes the preview itself, so the edit's render is already in the log.
        self.assertEqual(deck_calls, [(self.deck.name, [1, 2, 3, 4]), (self.deck.name, [3])])
        info = deck_preview._prepare_previews(self.deck, self.tmp / "c", Presentation(str(self.deck)))
        self.assertEqual((info["hits"], info["misses"]), (4, 0))

    def test_warm_failure_falls_back_to_the_cold_chain(self):
        """A dead bridge must cost latency, never the preview."""
        deck_preview.render_warm.get_renderer = lambda **kw: _FakeRenderer(self.root, fail=True)
        rendered = []

        def fake_cold(pptx, outdir, width=None):
            rendered.append(width)
            outdir.mkdir(parents=True, exist_ok=True)
            for i in range(1, 4):
                (outdir / f"slide-{i}.png").write_bytes(b"cold")
            return 3

        deck_preview._render = fake_cold
        info = deck_preview._prepare_previews(self.deck, self.tmp / "d", Presentation(str(self.deck)))
        self.assertEqual(info["renderer"], "cold")
        self.assertEqual(rendered, [deck_preview.preview_width()])
        self.assertTrue(any("warm render failed" in w for w in info["warnings"]))
        self.assertEqual(len(list((self.tmp / "d").glob("slide-*.png"))), 3)
        # The cold render is cached too, so the retry after a bridge crash is free.
        cached = sum(1 for k in self._keys()[:3] if render_cache.has(self.root, k))
        self.assertEqual(cached, 3)

    def test_cache_off_uses_the_cold_chain_and_ignores_the_renderer(self):
        fake = _FakeRenderer(self.root)
        deck_preview.render_warm.get_renderer = lambda **kw: fake
        deck_preview._render = lambda pptx, outdir, width=None: (outdir.mkdir(parents=True, exist_ok=True),
                                                                 (outdir / "slide-1.png").write_bytes(b"c"))[1]
        os.environ["PPT_RENDER_CACHE"] = "0"
        try:
            info = deck_preview._prepare_previews(self.deck, self.tmp / "e", Presentation(str(self.deck)))
        finally:
            os.environ.pop("PPT_RENDER_CACHE", None)
        self.assertEqual(fake.calls, [])
        self.assertEqual(info["renderer"], "cold")

    def test_preview_width_is_configurable_and_defaults_high(self):
        self.assertEqual(deck_preview.preview_width(), 1600)
        os.environ["PPT_PREVIEW_WIDTH"] = "1200"
        try:
            self.assertEqual(deck_preview.preview_width(), 1200)
        finally:
            os.environ.pop("PPT_PREVIEW_WIDTH", None)

    def test_prune_keeps_the_newest_renders(self):
        now = time.time()
        for i in range(4):
            png = self.tmp / f"p{i}.png"
            png.write_bytes(b"x")
            key = f"prunetest{i}0" * 4
            render_cache.store(self.root, key, png)
            os.utime(render_cache.path_for(self.root, key), (now + i, now + i))
        removed = render_cache.prune(self.root, limit=2)
        self.assertGreaterEqual(removed, 2)
        survivors = {k for k in (f"prunetest{i}0" * 4 for i in range(4))
                     if render_cache.has(self.root, k)}
        self.assertEqual(survivors, {f"prunetest{i}0" * 4 for i in (2, 3)})


class WarmQueueRegressions(unittest.TestCase):
    """The queue must coalesce duplicates and refuse work instead of growing unbounded."""

    def _renderer(self, queue_max=1):
        renderer = render_warm.WarmRenderer(python="unused", soffice_bin="unused",
                                            width=1600, queue_max=queue_max, timeout=5)
        self.addCleanup(renderer.close)
        return renderer

    def test_identical_requests_are_coalesced(self):
        renderer = self._renderer(queue_max=8)
        gate, calls = threading.Event(), []

        def slow_job(src, wanted, root):
            calls.append(sorted(wanted))
            gate.wait(10)
            return {i: root / f"x{i}.png" for i in wanted}
        renderer._render_job = slow_job
        deck = _TEMPLATES / "corporate-2slide.pptx"
        results = []
        threads = [threading.Thread(target=lambda: results.append(
            renderer.render(deck, {1: "key1", 2: "key2"}, _TMP))) for _ in range(4)]
        for t in threads:
            t.start()
        time.sleep(0.2)
        gate.set()
        for t in threads:
            t.join(20)
        self.assertEqual(len(calls), 1, "four identical previews must be one render")
        self.assertEqual(len(results), 4)
        self.assertGreaterEqual(renderer.stats()["coalesced"], 3)

    def test_full_queue_raises_instead_of_blocking(self):
        renderer = self._renderer(queue_max=1)
        gate = threading.Event()

        def stuck_job(src, wanted, root):
            gate.wait(10)
            return {}

        renderer._render_job = stuck_job
        deck = _TEMPLATES / "corporate-2slide.pptx"
        threading.Thread(target=lambda: renderer.render(deck, {1: "a"}, _TMP), daemon=True).start()
        deadline = time.time() + 5
        while time.time() < deadline and renderer.stats()["inflight"] == 0:
            time.sleep(0.02)  # wait until the worker is busy, so the queue is the only space
        threading.Thread(target=lambda: renderer.render(deck, {2: "b"}, _TMP), daemon=True).start()
        deadline = time.time() + 5
        while time.time() < deadline and renderer.stats()["queued"] == 0:
            time.sleep(0.02)
        with self.assertRaises(render_warm.RenderBusy):
            renderer.render(deck, {3: "c"}, _TMP)
        gate.set()

    def test_uno_python_detection_rejects_an_interpreter_without_pyuno(self):
        self.assertIsNone(render_warm.uno_python("/nonexistent/python3"))

    def test_engine_tag_records_the_libreoffice_version(self):
        """A LibreOffice upgrade renders differently, so old cached PNGs must not be reused."""
        render_warm._ENGINE_CACHE.clear()
        original = render_warm.soffice_version
        render_warm.soffice_version = lambda: "LibreOffice 99.1"
        try:
            self.assertIn("LibreOffice 99.1", render_warm.engine_tag())
        finally:
            render_warm.soffice_version = original
            render_warm._ENGINE_CACHE.clear()


if __name__ == "__main__":
    unittest.main(verbosity=2)
