"""RADIA: Regressions from the MCP deck alignment review."""
from __future__ import annotations

import os
import re
import sys
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _env import TMP  # noqa: E402
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
import auth  # noqa: E402
import deck_preview  # noqa: E402
import deck_registry  # noqa: E402
import paths  # noqa: E402
import server  # noqa: E402
import deck_edit_tools  # noqa: E402


class ReviewFixesTest(unittest.TestCase):
    def setUp(self):
        self.source = TMP / "out" / f"review-private-{uuid.uuid4().hex}.pptx"
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(self.source)
        self.tools = {t.name: t.fn for t in server.app._tool_manager.list_tools()}

    def test_paths_cannot_be_used_to_read_or_reassign_another_users_deck(self):
        with patch.dict(os.environ, {"PPT_AUTH_MODE": "token", "PPT_OWNERSHIP": "true"}):
            token = auth.set_current(auth.Identity(subject="alice", mode="token"))
            try:
                deck_id = deck_preview.remember(self.source)
            finally:
                auth.reset_current(token)
            token = auth.set_current(auth.Identity(subject="bob", mode="token"))
            try:
                with self.assertRaises(FileNotFoundError):
                    self.tools["get_deck_outline"](file=str(self.source))
                with patch.object(deck_preview, "_prepare_previews", return_value={"renderer": "none"}), \
                     patch.object(deck_preview, "_template_gallery", return_value=[]):
                    with self.assertRaises(FileNotFoundError):
                        self.tools["preview_deck"](file=str(self.source))
            finally:
                auth.reset_current(token)
            self.assertEqual(deck_registry.owner_for(deck_preview.preview_root(), deck_id),
                             auth.Identity(subject="alice", mode="token").scope_hint)

    def test_deck_paths_are_rejected_outside_the_upload_import_tool(self):
        with self.assertRaises(FileNotFoundError):
            deck_preview.resolve_ref(str(self.source))

    def test_import_only_accepts_chat_uploads_and_returns_an_owned_deck(self):
        uploaded = TMP / "uploads" / "user" / f"upload-{uuid.uuid4().hex}.pptx"
        Presentation().save(uploaded)
        tool = self.tools["import_deck"]
        with patch.dict(os.environ, {"PPT_AUTH_MODE": "token", "PPT_OWNERSHIP": "true"}), \
             patch.object(deck_preview, "_prepare_previews", return_value={"renderer": "none"}), \
             patch.object(deck_preview, "_template_gallery", return_value=[]):
            alice = auth.set_current(auth.Identity(subject="alice", mode="token"))
            try:
                result = tool(file=str(uploaded))
            finally:
                auth.reset_current(alice)
            bob = auth.set_current(auth.Identity(subject="bob", mode="token"))
            try:
                with self.assertRaises(FileNotFoundError):
                    deck_preview.resolve_ref(result["deck_id"])
                with self.assertRaises(FileNotFoundError):
                    tool(file=str(self.source))
            finally:
                auth.reset_current(bob)

    def test_image_ids_are_owner_scoped(self):
        with patch.dict(os.environ, {"PPT_AUTH_MODE": "token", "PPT_OWNERSHIP": "true"}):
            alice = auth.set_current(auth.Identity(subject="alice", mode="token"))
            try:
                image_id = deck_registry.new_image_id()
                image_path = TMP / "out" / f"image-{image_id}.png"
                Image.new("RGB", (8, 8), "red").save(image_path)
                deck_registry.remember_image(deck_preview.preview_root(), image_id, image_path, auth.current_identity().scope_hint)
                self.tools["enhance_image"](image_id=image_id)
            finally:
                auth.reset_current(alice)
            bob = auth.set_current(auth.Identity(subject="bob", mode="token"))
            try:
                with self.assertRaises(FileNotFoundError):
                    self.tools["enhance_image"](image_id=image_id)
                with self.assertRaises(ValueError):
                    paths.resolve_image(str(image_path))
            finally:
                auth.reset_current(bob)

    def test_template_apply_surfaces_element_failures_and_autosize_uses_enum(self):
        deck_id = deck_preview.remember(self.source)
        result = self.tools["create_slide_from_template"](
            file=deck_id, template_id="title_slide", content_mapping={"title": "New title"})
        self.assertNotIn("element_errors", result)
        texts = [shape.text for shape in Presentation(str(deck_preview.resolve_ref(result["deck_id"])))
                 .slides[1].shapes if shape.has_text_frame]
        self.assertTrue(any("New title" in text for text in texts))

    def test_template_element_errors_are_returned_without_internal_error_text(self):
        deck_id = deck_preview.remember(self.source)

        class Tool:
            fn = lambda self, **kwargs: {"slide_index": 0, "template_applied": {
                "elements_created": [{"role": "title", "error": "D:/private/template.json"}]}}

        class Manager:
            def get_tool(self, name):
                return Tool()

        class FakeApp:
            _tool_manager = Manager()

        with patch.object(deck_edit_tools, "_operation_app", return_value=FakeApp()):
            result = self.tools["apply_slide_template"](file=deck_id, slide_index=1,
                template_id="title_slide", content_mapping={"title": "x"})
        self.assertEqual(result["element_errors"], [{"role": "title", "message": "Template element could not be applied."}])

    def test_empty_preview_is_not_silently_ignored_and_hides_source_name(self):
        with patch.object(deck_preview, "_prepare_previews", return_value={"renderer": "none"}), \
             patch.object(deck_preview, "_template_gallery", return_value=[]):
            data = deck_preview.build_preview(deck_preview.remember(self.source))
        self.assertEqual(data["slides"], [])
        self.assertTrue(data["preview_error"])
        self.assertNotIn("source_name", data)
        self.assertFalse(re.search(r"[A-Za-z]:\\|/", data.get("preview_error") or ""))
        js = (Path(__file__).resolve().parents[1] / "apps/deck_studio/app.js").read_text(encoding="utf-8")
        self.assertNotRegex(js, r"data\.slides\.length\s*===\s*0\)\s*return")
        self.assertNotIn("Edit deck ${currentDeck}", js)

    def test_old_preview_metadata_does_not_reexpose_source_names(self):
        deck_id = "d" + uuid.uuid4().hex[:16]
        directory = deck_preview.preview_root() / deck_id
        directory.mkdir(parents=True)
        deck_registry.record(deck_preview.preview_root(), deck_id, self.source)
        (directory / "deck.json").write_text('{"deck_id":"' + deck_id + '","source_name":"private.pptx"}',
                                               encoding="utf-8")
        self.assertNotIn("source_name", deck_preview.last_preview())


if __name__ == "__main__":
    unittest.main()
