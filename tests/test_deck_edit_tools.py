"""RADIA: Editing uses authorized deck IDs and preserves the source deck."""
import asyncio
import os
import uuid
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _env import TMP  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "upstream"))

from pptx import Presentation  # noqa: E402
from PIL import Image  # noqa: E402
import auth  # noqa: E402
import deck_preview  # noqa: E402
import deck_registry  # noqa: E402
import paths  # noqa: E402
import server  # noqa: E402


class DeckEditToolsTest(unittest.TestCase):
    def setUp(self):
        self.source = TMP / "out" / "edit-source.pptx"
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        presentation.save(str(self.source))
        self.tools = {tool.name: tool.fn for tool in server.app._tool_manager.list_tools()}

    def test_edit_uses_deck_id_and_preserves_source(self):
        deck_id = deck_preview.remember(self.source)
        result = self.tools["add_shape"](file=deck_id, slide_index=1, shape_type="rectangle",
                                         left=1, top=1, width=2, height=1, text="Hello")
        self.assertNotEqual(result["deck_id"], deck_id)
        self.assertEqual(len(Presentation(str(self.source)).slides[0].shapes), 0)
        edited = Presentation(str(deck_preview.resolve_ref(result["deck_id"])))
        self.assertEqual(edited.slides[0].shapes[0].text, "Hello")
        info = self.tools["get_slide_info"](file=result["deck_id"], slide_index=1)
        self.assertEqual(info["geometry_units"], "inches")
        self.assertEqual(info["shapes"][0]["width"], 2.0)
        extracted = self.tools["extract_presentation_text"](file=result["deck_id"])
        self.assertEqual(extracted["slides_text"][0]["slide_index"], 1)
        second = self.tools["add_shape"](file=result["deck_id"], slide_index=1, shape_type="oval",
                                         left=3, top=1, width=1, height=1)
        self.assertEqual(len(Presentation(str(deck_preview.resolve_ref(second["deck_id"]))).slides[0].shapes), 2)
        with patch.object(deck_preview, "cache_enabled", return_value=False), \
             patch.object(deck_preview, "_render", return_value=0), \
             patch.object(deck_preview, "_template_gallery", return_value=[]):
            self.assertEqual(self.tools["preview_deck"](file=second["deck_id"])["deck_id"], second["deck_id"])
        with self.assertRaises(ValueError):
            self.tools["add_shape"](file=str(self.source), slide_index=1, shape_type="rectangle",
                                    left=1, top=1, width=2, height=1)

    def test_another_owner_cannot_edit_a_deck(self):
        with patch.dict(os.environ, {"PPT_AUTH_MODE": "token", "PPT_OWNERSHIP": "true"}):
            token = auth.set_current(auth.Identity(subject="alice", mode="token"))
            try:
                deck_id = deck_preview.remember(self.source)
                edited_id = self.tools["add_shape"](file=deck_id, slide_index=1, shape_type="rectangle",
                                                    left=1, top=1, width=1, height=1)["deck_id"]
            finally:
                auth.reset_current(token)
            token = auth.set_current(auth.Identity(subject="bob", mode="token"))
            try:
                for candidate in (deck_id, edited_id):
                    with self.subTest(deck=candidate), self.assertRaises(FileNotFoundError):
                        self.tools["get_slide_info"](file=candidate, slide_index=1)
            finally:
                auth.reset_current(token)

    def test_image_operations_use_image_ids(self):
        deck_id = deck_preview.remember(self.source)
        image_id = deck_registry.new_image_id()
        Image.new("RGB", (10, 10), "red").save(TMP / "out" / f"image-{image_id}.png")
        enhanced = self.tools["enhance_image"](image_id=image_id, brightness=0.5)
        self.assertNotEqual(enhanced["image_id"], image_id)
        self.assertTrue(paths.resolve_image(enhanced["image_id"]).is_file())
        result = self.tools["add_image"](file=deck_id, slide_index=1, image_id=enhanced["image_id"],
                                         left=1, top=1, width=1, height=1)
        edited = Presentation(str(deck_preview.resolve_ref(result["deck_id"])))
        self.assertEqual(len(edited.slides[0].shapes), 1)
        with self.assertRaises(ValueError):
            self.tools["add_image"](file=deck_id, slide_index=1, image_id="C:/secret.png",
                                    left=1, top=1, width=1, height=1)

    def test_new_tool_schema_does_not_expose_legacy_handles_or_paths(self):
        for name in ("add_shape", "manage_text", "apply_slide_template"):
            schema = server.app._tool_manager.get_tool(name).parameters
            self.assertIn("file", schema["required"])
            self.assertNotIn("presentation_id", schema["properties"])
            self.assertNotIn("image_paths", schema["properties"])
        self.assertIn("images", server.app._tool_manager.get_tool("apply_slide_template").parameters["properties"])
        self.assertNotIn("background_type", server.app._tool_manager.get_tool("add_slide").parameters["properties"])
        self.assertNotIn("apply_professional_design", self.tools)

    def test_edit_tool_runs_through_mcp_argument_validation(self):
        deck_id = deck_preview.remember(self.source)
        result = asyncio.run(server.app.call_tool("add_shape", {
            "file": deck_id, "slide_index": 1, "shape_type": "rectangle",
            "left": 1, "top": 1, "width": 1, "height": 1,
        }))
        self.assertNotEqual(result[1]["result"]["deck_id"], deck_id)

    def test_slide_templates_use_image_ids_and_preserve_existing_slides(self):
        deck_id = deck_preview.remember(self.source)
        result = self.tools["create_slide_from_template"](file=deck_id, template_id="title_slide",
                                                          content_mapping={"title": "New"})
        self.assertEqual(len(Presentation(str(deck_preview.resolve_ref(result["deck_id"]))).slides), 2)
        with self.assertRaises(ValueError):
            self.tools["apply_slide_template"](file=deck_id, slide_index=1, template_id="text_with_image",
                                                images={"supporting": "C:/secret.png"})

    def test_unimplemented_picture_effects_do_not_claim_success(self):
        deck_id = deck_preview.remember(self.source)
        with self.assertRaisesRegex(ValueError, "Only rotation"):
            self.tools["apply_picture_effects"](file=deck_id, slide_index=1, shape_index=0,
                                                effects={"shadow": {}})

    def _chart_deck(self):
        built = self.tools["build_deck"](
            template="corporate-2slide.pptx", title="Charts",
            slides=[{"title": "Intro", "bullets": ["a"]},
                    {"title": "Revenue", "chart": {"type": "column", "categories": ["Q1", "Q2"],
                                                      "series": [{"name": "2025", "values": [1, 2]}]}}])
        return built["deck_id"]

    def test_chart_updates_find_the_chart_without_a_shape_index(self):
        """The outline gives slide numbers but not shape indexes, so guessing one used to fail with
        "Shape is not a chart" and the model reported it could not find the graph."""
        deck_id = self._chart_deck()
        self.assertNotIn("shape_index",
                         server.app._tool_manager.get_tool("update_chart_data").parameters["required"])
        result = self.tools["update_chart_data"](file=deck_id, slide_index=3, categories=["Q1", "Q2"],
                                                 series_data=[{"name": "2026", "values": [7, 8]}])
        self.assertEqual(result["series_names"], ["2026"])
        chart = Presentation(str(deck_preview.resolve_ref(result["deck_id"]))).slides[2].shapes[1].chart
        self.assertEqual(list(chart.plots[0].series[0].values), [7.0, 8.0])

    def test_a_wrong_chart_reference_says_where_the_chart_is(self):
        deck_id = self._chart_deck()
        with self.assertRaisesRegex(ValueError, r"slide 3, shape 1"):
            self.tools["update_chart_data"](file=deck_id, slide_index=2, categories=["Q1"],
                                             series_data=[{"name": "x", "values": [1]}])
        with self.assertRaisesRegex(ValueError, r"shape 0.*slide 3, shape 1"):
            self.tools["update_chart_data"](file=deck_id, slide_index=3, shape_index=0,
                                             categories=["Q1"], series_data=[{"name": "x", "values": [1]}])
        with self.assertRaisesRegex(ValueError, "no charts"):
            # A file of its own: the id is derived from the path, so reusing self.source here would
            # hand later tests an already-registered (different-owner) entry.
            plain_path = TMP / "out" / f"no-charts-{uuid.uuid4().hex}.pptx"
            plain_path.write_bytes(self.source.read_bytes())
            plain = deck_preview.remember(plain_path)
            self.tools["update_chart_data"](file=plain, slide_index=1, categories=["Q1"],
                                             series_data=[{"name": "x", "values": [1]}])


if __name__ == "__main__":
    unittest.main()


class VerifyDensityTest(unittest.TestCase):
    """A 6th bullet shipped twice in the 2026-10-27 corpus: build_deck accepts it and nothing said so."""

    def setUp(self):
        self.tools = {tool.name: tool.fn for tool in server.app._tool_manager.list_tools()}

    def _build(self, bullets):
        return self.tools["build_deck"](template="corporate-2slide.pptx", title="Density",
                                        slides=[{"title": "One", "bullets": bullets}])["deck_id"]

    def test_six_bullets_are_reported_as_a_problem(self):
        data = self.tools["verify_deck"](file=self._build([f"b{i}" for i in range(6)]), expected_slides=2)
        self.assertFalse(data["ok"])
        self.assertTrue(any("6 bullets" in problem for problem in data["problems"]), data["problems"])

    def test_five_bullets_are_fine(self):
        data = self.tools["verify_deck"](file=self._build([f"b{i}" for i in range(5)]), expected_slides=2)
        self.assertTrue(data["ok"], data["problems"])
