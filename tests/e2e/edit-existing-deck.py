#!/usr/bin/env python
"""Ad-hoc E2E: does the MODEL edit an existing deck through the id-only tools?

The 108-case corpus only builds decks from scratch, and the browser E2E opens the style
panel without applying it. This drives one LibreChat conversation through both edit shapes:

    turn 1  build + preview a 3-slide deck
    turn 2  "make slide 2's title bold and 40pt"            -> style_slide, in place, same id
    turn 3  "add a slide titled Next steps with 2 bullets"  -> NEW id, previous deck intact
            (the skill's default is a rebuild; a vendored edit tool is equally valid)

Contract under test (docs/decisions.md, 2026-10-27):
  * style_slide / style_text restyle in place and keep the deck id; they re-render the previews;
  * a structural edit yields a NEW deck id and leaves the deck it started from alone.
"""
from __future__ import annotations

import json
import sys
import uuid
from pathlib import Path

ROOT = Path("/home/hyma/dev/ppt-mcp")
sys.path.insert(0, str(ROOT / "tests"))
sys.path.insert(0, str(ROOT))

import run_suite as R  # noqa: E402
from pptx import Presentation  # noqa: E402

ok = True
VENDORED_EDITS = {"add_slide", "manage_text", "populate_placeholder", "add_bullet_points",
                  "add_table", "add_shape", "add_chart"}


def say(name, passed, detail=""):
    global ok
    ok = ok and passed
    print(f"{'PASS' if passed else 'FAIL'}  {name}{' — ' + detail if detail else ''}", flush=True)


def turn(token, agent_id, text, conversation_id=None, parent=None):
    conv = R.post("/api/agents/chat", {
        "text": text,
        "conversationId": conversation_id,
        "parentMessageId": parent or "00000000-0000-0000-0000-000000000000",
        "messageId": str(uuid.uuid4()),
        "endpoint": "agents", "endpointType": "agents", "agent_id": agent_id,
    }, token)
    events = R.stream(token, conv["streamId"], conv["generationCreatedAt"])
    final = next((e for e in events if e.get("final") and "responseMessage" in e), {})
    calls, reply = R.parse_tool_calls(events)
    return conv, final.get("responseMessage", {}), calls, reply


def payload(output):
    try:
        value = R.unwrap(json.loads(output))
        return value if isinstance(value, dict) else {}
    except Exception:
        return {}


def tool_base(name):
    return name.replace("_mcp_ppt", "").replace("mcp_ppt_", "")


def title(path):
    """(text, size_pt, bold) of the first non-empty run on slide 2."""
    slide = Presentation(path).slides[1]
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.text.strip():
                        return (shape.text_frame.text.strip().split("\n")[0],
                                run.font.size.pt if run.font.size else None, run.font.bold)
    return None, None, None


def main():
    token = R.login("hyma@local.test", "PiTest!2026")
    agent_id = json.loads(Path("/tmp/agent.json").read_text())["id"]

    conv, first, calls, _ = turn(
        token, agent_id, "Make a 3-slide deck about our Q3 results and show it to me.")
    # The agent may rebuild once after verify_deck; the deck it *showed* is the last one previewed.
    ids = [payload(c["output"]).get("deck_id") for c in calls
           if tool_base(c["name"]) in ("build_deck", "preview_deck")]
    deck_id = ids[-1] if ids else None
    print(f"turn 1 tools: {[tool_base(c['name']) for c in calls]}", flush=True)
    say("turn 1 built a deck", bool(deck_id), deck_id or "no deck_id")
    if not deck_id:
        return 1
    source = R._deck_file(deck_id)
    say("turn 1 previewed it", any(tool_base(c["name"]) == "preview_deck" for c in calls))

    conv2, second, calls2, _ = turn(
        token, agent_id,
        "Make the title on slide 2 bold and 40 point.",
        conversation_id=conv["conversationId"], parent=first.get("messageId"))
    edits = [payload(c["output"]) for c in calls2 if tool_base(c["name"]) == "style_slide"]
    print(f"turn 2 tools: {[tool_base(c['name']) for c in calls2]}", flush=True)
    say("the model called style_slide", bool(edits), json.dumps(edits[-1])[:120] if edits else "none")
    say("style_slide kept the deck id (in place)",
        bool(edits) and edits[-1].get("deck_id") == deck_id, str(edits[-1].get("deck_id") if edits else None))
    say("style_slide returned fresh slide URLs", bool(edits) and bool(edits[-1].get("slides")),
        f"{len(edits[-1].get('slides') or [])} slides" if edits else "none")
    text, size, bold = title(source)
    say("slide 2 title is 40 pt", size is not None and round(size) == 40, f"{size} pt")
    say("slide 2 title is bold", bold is True, str(bold))

    conv3, third, calls3, reply3 = turn(
        token, agent_id,
        "Add a new slide at the end titled 'Next steps' with two short bullets, then show me the deck.",
        conversation_id=conv["conversationId"], parent=second.get("messageId"))
    print(f"turn 3 tools: {[tool_base(c['name']) for c in calls3]}", flush=True)
    structural = [c for c in calls3 if tool_base(c["name"]) in VENDORED_EDITS]
    rebuilt = [payload(c["output"]).get("deck_id") for c in calls3
               if tool_base(c["name"]) == "build_deck"]
    # The skill tells the model to rebuild for a wording/structural change and to use the vendored
    # tools for repair, so either answer is legitimate. Record which, then judge the artefact.
    path = ("vendored:" + ",".join(tool_base(c["name"]) for c in structural)) if structural \
        else "rebuild (build_deck)" if rebuilt else "no edit"
    print(f"      edit path chosen: {path}", flush=True)
    final_id = (payload(c["output"]).get("deck_id") for c in reversed(calls3)
                if tool_base(c["name"]) in ("build_deck", "preview_deck"))
    final_id = next((i for i in final_id if i), None)
    say("the model produced an updated deck", bool(final_id) and final_id != deck_id,
        f"{deck_id} -> {final_id}")
    if final_id and final_id != deck_id:
        edited = R._deck_file(final_id)
        say("the edited deck has 4 slides", len(Presentation(edited).slides) == 4,
            f"{len(Presentation(edited).slides)} slides")
        say("the previous deck still has 3 slides", len(Presentation(source).slides) == 3,
            f"{len(Presentation(source).slides)} slides")
        say("the new slide is titled Next steps",
            any("next steps" in (sh.text_frame.text or "").lower()
                for sh in Presentation(edited).slides[3].shapes if sh.has_text_frame))
    print(f"\nconversation: {R.BASE}/c/{conv['conversationId']}")
    print(f"reply: {(reply3 or '').strip()[:200]}")
    print("ALL PASS" if ok else "SOME CHECKS FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
