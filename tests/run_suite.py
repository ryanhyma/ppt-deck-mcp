#!/usr/bin/env python
"""Run the deck corpus against LibreChat + the local 27B, one conversation per case.

    .venv/bin/python tests/run_suite.py --limit 12
    .venv/bin/python tests/run_suite.py --category charts --limit 20
    .venv/bin/python tests/run_suite.py --ids chart-01,table-03

Writes tests/results/<stamp>.json (full evidence) and <stamp>.md (readable report).
Each case opens its OWN conversation, so the chat list can be reviewed later.
"""
from __future__ import annotations

import argparse
import json
import statistics
import sys
import time
import urllib.error
import urllib.request
import uuid
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pptx import Presentation  # noqa: E402

BASE = "http://localhost:3080"
HERE = Path(__file__).resolve().parent
HDRS = {
    "Content-Type": "application/json",
    "Origin": BASE,
    "User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/156.0 Safari/537.36",
}


class AuthError(Exception):
    """Session token expired (SESSION_EXPIRY is 15 min on this deployment)."""


def post(path: str, payload: dict, token: str | None = None) -> dict:
    req = urllib.request.Request(BASE + path, data=json.dumps(payload).encode(), headers={
        **HDRS, **({"Authorization": f"Bearer {token}"} if token else {})})
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            return json.loads(r.read().decode() or "{}")
    except urllib.error.HTTPError as exc:
        if exc.code == 401:
            raise AuthError(path) from exc
        raise


def login(email: str, password: str) -> str:
    return post("/api/auth/login", {"email": email, "password": password})["token"]


def stream(token: str, stream_id: str, created_at: int) -> list[dict]:
    url = f"{BASE}/api/agents/chat/stream/{stream_id}?generationCreatedAt={created_at}"
    req = urllib.request.Request(url, headers={**HDRS, "Authorization": f"Bearer {token}",
                                               "Accept": "text/event-stream"})
    events: list[dict] = []
    try:
        handle = urllib.request.urlopen(req, timeout=600)
    except urllib.error.HTTPError as exc:
        if exc.code == 401:
            raise AuthError("stream") from exc
        raise
    with handle as r:
        for raw in r:
            line = raw.decode("utf-8", "replace").strip()
            if not line.startswith("data: "):
                continue
            try:
                events.append(json.loads(line[6:]))
            except Exception:
                continue
            if events and events[-1].get("final") is True and "responseMessage" in events[-1]:
                break
    return events


def unwrap(value):
    while isinstance(value, dict) and set(value) == {"result"}:
        value = value["result"]
    return value


def parse_tool_calls(events: list[dict]) -> tuple[list[dict], str]:
    calls, text = [], ""
    for e in events:
        if e.get("event") == "on_run_step_completed":
            r = e.get("data", {}).get("result", {})
            tc = r.get("tool_call") or {}
            if tc.get("name"):
                calls.append({"name": tc["name"], "output": tc.get("output", "")})
        if e.get("final"):
            for c in (e.get("responseMessage") or {}).get("content") or []:
                if c.get("type") == "text":
                    text += c.get("text", "")
    return calls, text


def _deck_file(deck_id: str | None) -> str | None:
    """The MCP surface returns an opaque deck id; this runner resolves it on the host to inspect it."""
    if not deck_id:
        return None
    try:
        from deck_preview import resolve_ref  # noqa: PLC0415 - host-side helper, not part of the MCP API

        return str(resolve_ref(deck_id))
    except Exception:
        return None


def deck_from_calls(calls: list[dict]) -> tuple[str | None, str | None]:
    """(file, deck_id) from preview_deck/build_deck outputs, preferring preview_deck."""
    found = {"preview": (None, None), "build": (None, None)}
    for c in calls:
        try:
            payload = unwrap(json.loads(c["output"]))
        except Exception:
            continue
        if not isinstance(payload, dict):
            continue
        if c["name"].startswith("preview_deck"):
            found["preview"] = (_deck_file(payload.get("deck_id")), payload.get("deck_id"))
        elif c["name"].startswith("build_deck"):
            found["build"] = (_deck_file(payload.get("deck_id")), payload.get("deck_id"))
    return found["preview"] if found["preview"][0] else found["build"]


def inspect_deck(path: str) -> dict:
    prs = Presentation(path)
    charts = tables = pictures = 0
    bullets, titles, overflow = [], [], []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_chart:
                charts += 1
            elif shape.has_table:
                tables += 1
            elif shape.shape_type == 13:  # picture
                pictures += 1
        for ph in slide.placeholders:
            if not ph.has_text_frame or not ph.text_frame.text.strip():
                continue
            text = ph.text_frame.text
            if ph.placeholder_format.idx == 0:
                titles.append(text.strip().split("\n")[0])
            elif ph.placeholder_format.idx == 1:
                rows = [r.strip() for r in text.split("\n") if r.strip()]
                bullets.extend(rows)
                if len(rows) > 5:
                    overflow.append(len(rows))
    return {
        "slides": len(prs.slides),
        "charts": charts,
        "tables": tables,
        "pictures": pictures,
        "bullets": len(bullets),
        "max_bullet_chars": max((len(b) for b in bullets), default=0),
        "titles": titles,
        "bullet_overflow_slides": overflow,
    }


def is_tool_error(output: str) -> bool:
    head = (output or "").strip()[:80].lower()
    return ("error executing tool" in head or head.startswith("error")
            or '"iserror": true' in (output or "").lower())


def judge(case: dict, got: dict, calls: list[dict], error: str | None) -> tuple[bool, list[str], list[str]]:
    """Returns (output_ok, failures, process_warnings).

    pass/fail is about the artefact — did the deck match the request. Tool errors and a missing
    preview_deck are process signals: the agent may have failed once, corrected itself, and still
    produced the right file, which is a different (and less interesting) problem.
    """
    exp, fails = case["expect"], []
    if error:
        fails.append(f"error: {error[:120]}")
    if got is None:
        fails.append("no deck produced")
        return False, fails, []
    if exp["slides"] and got["slides"] != exp["slides"]:
        fails.append(f"slides {got['slides']}≠{exp['slides']}")
    if exp["chart"] and got["charts"] < 1:
        fails.append("chart missing")
    if exp["table"] and got["tables"] < 1:
        fails.append("table missing")
    if exp["image"] and got["pictures"] < 1:
        fails.append("image missing")
    if got["max_bullet_chars"] > 110:
        fails.append(f"bullet too long ({got['max_bullet_chars']})")
    if got["bullet_overflow_slides"]:
        fails.append(f"{len(got['bullet_overflow_slides'])} slides >5 bullets")
    warnings = []
    errs = [c["name"] for c in calls if is_tool_error(c.get("output") or "")]
    if errs:
        warnings.append(f"tool errors: {','.join(sorted(set(errs)))}")
    if not any(c["name"].startswith("preview_deck") for c in calls):
        warnings.append("preview_deck not called")
    return (not fails), fails, warnings


def run_case(case: dict, token: str, agent_id: str) -> dict:
    started = time.time()
    conv = post("/api/agents/chat", {
        "text": case["prompt"],
        "conversationId": None,
        "parentMessageId": "00000000-0000-0000-0000-000000000000",
        "messageId": str(uuid.uuid4()),
        "endpoint": "agents",
        "endpointType": "agents",
        "agent_id": agent_id,
    }, token)
    events = stream(token, conv["streamId"], conv["generationCreatedAt"])
    calls, text = parse_tool_calls(events)
    file, deck_id = deck_from_calls(calls)
    got = inspect_deck(file) if file and Path(file).is_file() else None
    err = None
    if not file:
        err = "no build_deck/preview_deck result"
    ok, fails, warnings = judge(case, got, calls, err)
    return {
        "id": case["id"],
        "category": case["category"],
        "note": case.get("note", ""),
        "prompt": case["prompt"],
        "expect": case["expect"],
        "pass": ok,
        "failures": fails,
        "warnings": warnings,
        "conversationId": conv["conversationId"],
        "url": f"{BASE}/c/{conv['conversationId']}",
        "file": file,
        "deck_id": deck_id,
        "tools": [c["name"] for c in calls],
        "tool_errors": [c["name"] for c in calls if is_tool_error(c.get("output") or "")],
        "got": got,
        "final_text": text[:400],
        "seconds": round(time.time() - started, 1),
    }


def report(results: list[dict], stamp: str) -> str:
    by_cat: dict[str, list[dict]] = {}
    for r in results:
        by_cat.setdefault(r["category"], []).append(r)
    lines = [f"# Deck suite — {stamp}", "",
             f"Cases: **{len(results)}** · passed **{sum(r['pass'] for r in results)}** · "
             f"median **{statistics.median([r['seconds'] for r in results]):.0f}s** per case", "",
             "| category | cases | pass | rate | median s |", "|---|---|---|---|---|"]
    for cat, rows in sorted(by_cat.items()):
        p = sum(r["pass"] for r in rows)
        med = statistics.median([r["seconds"] for r in rows])
        lines.append(f"| {cat} | {len(rows)} | {p} | {p/len(rows):.0%} | {med:.0f} |")
    lines += ["", "## Failures", ""]
    failures = [r for r in results if not r["pass"]]
    if not failures:
        lines.append("None.")
    for r in failures:
        lines.append(f"- **{r['id']}** ({r['category']}) — {'; '.join(r['failures'])}  \n"
                     f"  {r['prompt'][:120]}")
    lines += ["", "## Per case", "",
              "| id | cat | pass | slides (exp/got) | chart | table | img | s | tools |",
              "|---|---|---|---|---|---|---|---|---|"]
    for r in sorted(results, key=lambda x: (x["category"], x["id"])):
        g = r["got"] or {}
        lines.append(
            f"| {r['id']} | {r['category']} | {'✅' if r['pass'] else '❌'} | "
            f"{r['expect']['slides']}/{g.get('slides','-')} | "
            f"{g.get('charts','-')} | {g.get('tables','-')} | {g.get('pictures','-')} | "
            f"{r['seconds']:.0f} | {len(r['tools'])} |")
    return "\n".join(lines) + "\n"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--cases", default=str(HERE / "cases.json"))
    ap.add_argument("--limit", type=int, default=0, help="0 = all")
    ap.add_argument("--category", default="", help="comma-separated")
    ap.add_argument("--ids", default="", help="comma-separated case ids")
    ap.add_argument("--email", default="hyma@local.test")
    ap.add_argument("--password", default="PiTest!2026")
    ap.add_argument("--outdir", default=str(HERE / "results"))
    ap.add_argument("--resume", action="store_true",
                    help="skip cases that already passed in any results/*.json")
    args = ap.parse_args()

    corpus = json.loads(Path(args.cases).read_text())["cases"]
    if args.category:
        cats = {c.strip() for c in args.category.split(",")}
        corpus = [c for c in corpus if c["category"] in cats]
    if args.ids:
        ids = {i.strip() for i in args.ids.split(",")}
        corpus = [c for c in corpus if c["id"] in ids]

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    if args.resume:
        passed = set()
        for f in sorted(outdir.glob("*.json")):
            try:
                for r in json.loads(f.read_text()).get("results", []):
                    if r.get("pass"):
                        passed.add(r["id"])
            except Exception:
                continue
        before = len(corpus)
        corpus = [c for c in corpus if c["id"] not in passed]
        print(f"resume: {before - len(corpus)} already passed, {len(corpus)} to run", flush=True)
    if args.limit:
        corpus = corpus[: args.limit]

    token = login(args.email, args.password)
    agent_file = Path("/tmp/agent.json")
    if not agent_file.is_file():
        print("agent id missing (/tmp/agent.json)", file=sys.stderr)
        return 2
    agent_id = json.loads(agent_file.read_text())["id"]

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    results: list[dict] = []

    def flush() -> None:
        (outdir / f"{stamp}.json").write_text(json.dumps({"stamp": stamp, "results": results}, indent=1))
        (outdir / f"{stamp}.md").write_text(report(results, stamp))

    for i, case in enumerate(corpus, 1):
        print(f"[{i}/{len(corpus)}] {case['id']:<12} {case['prompt'][:70]}", flush=True)
        try:
            try:
                r = run_case(case, token, agent_id)
            except AuthError:
                print("      ~ session expired, re-authenticating", flush=True)
                token = login(args.email, args.password)
                r = run_case(case, token, agent_id)
        except Exception as exc:
            r = {"id": case["id"], "category": case["category"], "prompt": case["prompt"],
                 "expect": case["expect"], "pass": False, "failures": [f"runner: {exc}"],
                 "conversationId": None, "url": "", "file": None, "deck_id": None,
                 "tools": [], "tool_errors": [], "got": None, "final_text": "", "seconds": 0,
                 "warnings": [], "note": case.get("note", "")}
        results.append(r)
        flush()
        print(f"      -> {'PASS' if r['pass'] else 'FAIL: ' + '; '.join(r['failures'])}"
              f"  {r['seconds']}s  {r['url']}", flush=True)

    print("\n" + report(results, stamp).split("## Failures")[0])
    print(f"report: {outdir / (stamp + '.md')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
