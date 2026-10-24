"""Shared visual style: the deterministic chrome every generated slide gets.

Kept as code (not slide-master XML) so it works the same with python-pptx on Linux
and Windows and cannot be broken by template edits.
"""
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1F, 0x38, 0x64)
ACCENT = RGBColor(0x2A, 0x9D, 0x8F)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
INK = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x9A, 0xA0, 0xA6)
LIGHT = RGBColor(0xE9, 0xEC, 0xEF)

# Chart series palette: teal, navy, sand, coral, deep teal, sage.
CHART_PALETTE = [
    RGBColor(0x2A, 0x9D, 0x8F),
    RGBColor(0x1F, 0x38, 0x64),
    RGBColor(0xE9, 0xC4, 0x6A),
    RGBColor(0xE7, 0x6F, 0x51),
    RGBColor(0x26, 0x46, 0x53),
    RGBColor(0x8A, 0xB1, 0x7D),
]

WIDTH = Inches(13.333)
HEIGHT = Inches(7.5)

FOOTER = "Acme Corp  |  Confidential"


def _no_shadow(shape) -> None:
    shape.shadow.inherit = False


def decorate(slide) -> None:
    """Bottom rule + accent bar + footer wordmark. Purely additive, no placeholders."""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, HEIGHT - Inches(0.22), WIDTH, Inches(0.22))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    _no_shadow(bar)

    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, HEIGHT - Inches(0.26), WIDTH, Inches(0.04))
    rule.fill.solid()
    rule.fill.fore_color.rgb = NAVY
    rule.line.fill.background()
    _no_shadow(rule)

    foot = slide.shapes.add_textbox(Inches(0.55), HEIGHT - Inches(0.60), Inches(7), Inches(0.3))
    ftf = foot.text_frame
    ftf.text = FOOTER
    for p in ftf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(10)
            r.font.color.rgb = MUTED
            r.font.name = "Calibri"


def title_rule(slide, top: float = 1.28, width: float = 1.1) -> None:
    """Short teal underline beneath a content-slide title."""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.92), Inches(top), Inches(width), Inches(0.055)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    _no_shadow(bar)


def page_number(slide, number: int) -> None:
    box = slide.shapes.add_textbox(WIDTH - Inches(1.35), HEIGHT - Inches(0.62), Inches(0.8), Inches(0.3))
    tf = box.text_frame
    tf.text = f"{number:02d}"
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    for r in p.runs:
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = MUTED
        r.font.name = "Calibri"
