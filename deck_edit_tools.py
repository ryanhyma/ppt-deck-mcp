"""RADIA: Deck-ID based access to the distinct slide editing operations."""
import inspect
import re
import uuid

from mcp.server.fastmcp import FastMCP
from pptx import Presentation
from pptx.util import Inches

import deck_registry
import auth
import paths
import ppt_mcp_server as legacy
from deck_preview import preview_root, remember, resolve_ref
from deck_tools import _output_dir, _save_atomic

# These edit existing slides without rebuilding the deck from an incomplete outline.
_OPERATIONS = {
    "get_slide_info": ("content", False),
    "extract_slide_text": ("content", False),
    "extract_presentation_text": ("content", False),
    "get_presentation_info": ("presentation", False),
    "add_slide": ("content", True),
    "populate_placeholder": ("content", True),
    "add_bullet_points": ("content", True),
    "manage_text": ("content", True),
    "add_table": ("structural", True),
    "format_table_cell": ("structural", True),
    "add_shape": ("structural", True),
    "add_chart": ("structural", True),
    "update_chart_data": ("chart", True),
    "add_connector": ("connector", True),
    "manage_hyperlinks": ("hyperlink", True),
    "manage_slide_masters": ("master", False),
    "set_core_properties": ("presentation", True),
    "apply_picture_effects": ("professional", True),
    "optimize_slide_text": ("template", True),
    "apply_slide_template": ("template", True),
    "create_slide_from_template": ("template", True),
}


def _save_copy(presentation):
    destination = _output_dir() / f"deck-{uuid.uuid4().hex}.pptx"
    _save_atomic(presentation, destination)
    return remember(destination)


def _open_deck(deck_id):
    if not deck_registry.DECK_ID_RE.fullmatch(deck_id):
        raise ValueError("Use the deck ID returned by a deck tool.")
    return Presentation(str(resolve_ref(deck_id)))


def _image_path(image_id):
    if not re.fullmatch(r"i[0-9a-f]{16}", image_id):
        raise ValueError("Use an image ID returned by render_svg.")
    return str(paths.resolve_image(image_id))


def _operation_app(group, deck_id, presentation):
    """Bind upstream operations to one request's presentation, without shared global state."""
    app = FastMCP("deck-edit")
    presentations = {deck_id: presentation}
    current = lambda: deck_id
    validators = (legacy.validate_parameters, legacy.is_positive, legacy.is_non_negative,
                  legacy.is_in_range, legacy.is_valid_rgb)
    if group == "content":
        legacy.register_content_tools(app, presentations, current, *validators)
    elif group == "structural":
        legacy.register_structural_tools(app, presentations, current, *validators, legacy.add_shape_direct)
    elif group in {"chart", "connector", "hyperlink", "master"}:
        getattr(legacy, f"register_{group}_tools")(app, presentations, current, *validators)
    elif group == "presentation":
        legacy.register_presentation_tools(app, presentations, current, legacy.get_template_search_directories)
    elif group == "professional":
        legacy.register_professional_tools(app, presentations, current)
    elif group == "template":
        legacy.register_template_tools(app, presentations, current)
    return app


def register_edit_tools(app, upstream_tools):
    for name, (group, modifies) in _OPERATIONS.items():
        original = upstream_tools[name].fn
        signature = inspect.signature(original, eval_str=True)
        omitted = {"presentation_id"}
        if name == "add_slide":
            omitted.update({"background_type", "background_colors", "gradient_direction", "color_scheme"})
        params = [p.replace(name="images") if p.name == "image_paths" else p
                  for p in signature.parameters.values() if p.name not in omitted]
        params.append(inspect.Parameter("file", inspect.Parameter.KEYWORD_ONLY, annotation=str))

        def make_tool(tool_name, tool_group, writes, tool_signature, tool_doc):
            def edit(**kwargs):
                deck_id = kwargs.pop("file")
                presentation = _open_deck(deck_id)
                if tool_name == "apply_picture_effects" and set(kwargs.get("effects", {})) - {"rotation"}:
                    raise ValueError("Only rotation is supported for picture effects.")
                if "images" in kwargs:
                    kwargs["image_paths"] = {role: _image_path(image_id)
                                             for role, image_id in (kwargs.pop("images") or {}).items()}
                if kwargs.get("slide_index") is not None:
                    if not 1 <= kwargs["slide_index"] <= len(presentation.slides):
                        raise ValueError(f"slide_index must be 1..{len(presentation.slides)}")
                    kwargs["slide_index"] -= 1
                isolated = _operation_app(tool_group, deck_id, presentation)
                operation = isolated._tool_manager.get_tool(tool_name).fn
                result = operation(presentation_id=deck_id, **kwargs)
                if tool_name in {"apply_slide_template", "create_slide_from_template"}:
                    if not isinstance(result, dict) or result.get("error"):
                        return {"error": "The slide template could not be applied."}
                    failures = [item for item in result.get("template_applied", {}).get("elements_created", [])
                                if item.get("error")]
                    result = {key: result[key] for key in ("slide_index",) if key in result}
                    if failures:
                        result["element_errors"] = [
                            {"role": item.get("role", "unknown"), "message": "Template element could not be applied."}
                            for item in failures
                        ]
                if not isinstance(result, dict) or result.get("error"):
                    return result
                if writes:
                    deck_id = _save_copy(presentation)
                result.pop("presentation_id", None)
                if re.search(r"\bslide\s+\d+\b", str(result.get("message", "")), re.IGNORECASE):
                    result.pop("message", None)
                for key in ("slide_index", "new_slide_index"):
                    if isinstance(result.get(key), int):
                        result[key] += 1
                if tool_name == "extract_presentation_text":
                    for slide in result.get("slides_text", []):
                        if isinstance(slide.get("slide_index"), int):
                            slide["slide_index"] += 1
                if tool_name == "get_slide_info":
                    for shape in result.get("shapes", []):
                        for key in ("left", "top", "width", "height"):
                            if isinstance(shape.get(key), (int, float)):
                                shape[key] = round(shape[key] / 914400, 4)
                    result["geometry_units"] = "inches"
                if tool_name == "get_presentation_info":
                    for key in ("slide_width", "slide_height"):
                        if isinstance(result.get(key), (int, float)):
                            result[key] = round(result[key] / 914400, 4)
                    result["geometry_units"] = "inches"
                result["deck_id"] = deck_id
                return result

            edit.__name__ = tool_name
            edit.__signature__ = tool_signature.replace(parameters=params)
            doc_lines = [line.replace("0-based", "1-based") if "slide_index" in line else line
                         for line in tool_doc.splitlines()
                         if "presentation_id" not in line and "image_paths" not in line]
            edit.__doc__ = ("Use file as the deck ID; slide_index is 1-based and shape indexes are 0-based. "
                            "For edits, use the returned deck_id for subsequent calls.\n\n" + "\n".join(doc_lines))
            if tool_name in {"apply_slide_template", "create_slide_from_template"}:
                edit.__doc__ += "\nimages maps image roles to image IDs returned by render_svg."
            if tool_name == "apply_picture_effects":
                edit.__doc__ += "\nOnly the rotation effect is supported."
            if tool_name == "add_slide":
                edit.__doc__ = "Add a slide using a layout from the existing deck. Use the returned deck_id for later calls."
            if tool_name == "get_slide_info":
                edit.__doc__ += "\nShape positions and dimensions are in inches."
            return edit

        app.add_tool(make_tool(name, group, modifies, signature, inspect.getdoc(original) or ""))

    @app.tool()
    def add_image(file: str, slide_index: int, image_id: str, left: float, top: float,
                  width: float, height: float) -> dict:
        """Add a render_svg image to an existing slide. Use the returned deck_id for later calls.

        slide_index is 1-based; positions and dimensions are in inches.
        """
        if width <= 0 or height <= 0:
            raise ValueError("Image dimensions must be positive.")
        presentation = _open_deck(file)
        if not 1 <= slide_index <= len(presentation.slides):
            raise ValueError("slide_index is out of range.")
        slide = presentation.slides[slide_index - 1]
        slide.shapes.add_picture(_image_path(image_id), Inches(left), Inches(top),
                                 Inches(width), Inches(height))
        return {"deck_id": _save_copy(presentation), "slide_index": slide_index,
                "shape_index": len(slide.shapes) - 1}

    @app.tool()
    def enhance_image(image_id: str, brightness: float = 1.0, contrast: float = 1.0,
                      saturation: float = 1.0, sharpness: float = 1.0,
                      blur_radius: float = 0, filter_type: str = "") -> dict:
        """Enhance a generated image and return an image ID for use in a deck."""
        from PIL import Image, ImageEnhance, ImageFilter

        factors = (brightness, contrast, saturation, sharpness)
        if not all(0 <= value <= 5 for value in factors) or not 0 <= blur_radius <= 20:
            raise ValueError("Image adjustment values are out of range.")
        filters = {"BLUR": ImageFilter.BLUR, "SHARPEN": ImageFilter.SHARPEN,
                   "SMOOTH": ImageFilter.SMOOTH, "EDGE_ENHANCE": ImageFilter.EDGE_ENHANCE}
        if filter_type and filter_type.upper() not in filters:
            raise ValueError("Unsupported image filter.")
        with Image.open(_image_path(image_id)) as source:
            image = source.copy()
        for enhancer, factor in zip((ImageEnhance.Brightness, ImageEnhance.Contrast,
                                     ImageEnhance.Color, ImageEnhance.Sharpness), factors):
            if factor != 1:
                image = enhancer(image).enhance(factor)
        if blur_radius:
            image = image.filter(ImageFilter.GaussianBlur(radius=blur_radius))
        if filter_type:
            image = image.filter(filters[filter_type.upper()])
        new_id = deck_registry.new_image_id()
        destination = _output_dir() / f"image-{new_id}.png"
        image.save(destination, format="PNG")
        identity = auth.current_identity()
        deck_registry.remember_image(preview_root(), new_id, destination,
                                     identity.scope_hint if not identity.anonymous else "anonymous")
        return {"image_id": new_id, "width_px": image.width, "height_px": image.height}

    @app.tool()
    def list_slide_templates() -> dict:
        """List the available slide layout templates."""
        result = upstream_tools["list_slide_templates"].fn()
        result.pop("usage_examples", None)
        return result

    @app.tool()
    def get_template_info(template_id: str) -> dict:
        """Describe the elements in a slide layout template."""
        return upstream_tools["get_template_info"].fn(template_id=template_id)
